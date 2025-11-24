import os
import re
from docx import Document
from openpyxl import Workbook
from shutil import copy

# -------------------------------
# CONFIG
# -------------------------------
SOURCE_FOLDER = "source_docs"
PROCESSED_FOLDER = "processed"
DONE_FOLDER = "done"
TOTAL_SCORE = 60

# Option mapping A–E → Excel columns
OPTION_MAP = {"A": "C", "B": "D", "C": "E", "D": "F", "E": "G"}

# -------------------------------
# HELPERS
# -------------------------------
def extract_questions(doc):
    questions = []
    current_q = None

    # Get all non-empty lines from document
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for line in lines:
        # Detect inline options (e.g. "question ______ a) opt1 b) opt2 Answer: A")
        inline_opts = re.findall(r"\(?([a-eA-E])[).]([^(\n]+)", line)
        if inline_opts:
            # Save previous question
            if current_q:
                questions.append(current_q)
            # Extract question text before options
            parts = line.split(' a)')
            question_text = parts[0].strip() if len(parts) > 1 else line
            # Remove Answer: if present
            if ' Answer:' in question_text:
                question_text = question_text.split(' Answer:')[0].strip()
            current_q = {
                "sn": str(len(questions) + 1),
                "question": question_text,
                "options": {},
                "correct": None
            }
            # Add options
            for letter, option_text in inline_opts:
                letter = letter.upper()
                current_q["options"][letter] = option_text.strip()
            # Check for answer in the line
            ans_match = re.search(r"(?i)Answer[:\s]*([a-eA-E])", line)
            if ans_match:
                current_q["correct"] = ans_match.group(1).upper()
            continue

        # Detect question number line (e.g. "1.", "2)", "3 ")
        q_match = re.match(r"^(\d+)[\.\)]?\s+(.*)", line)
        if q_match:
            # Save previous question before starting new one
            if current_q:
                questions.append(current_q)
            current_q = {
                "sn": q_match.group(1),
                "question": q_match.group(2).strip(),
                "options": {},
                "correct": None
            }
            continue

        # Detect options (a), b), etc.)
        opt_match = re.match(r"^\(?([a-eA-E])[).]?\s*(.*)", line)
        if opt_match and current_q:
            letter = opt_match.group(1).upper()
            option_text = opt_match.group(2).strip()
            current_q["options"][letter] = option_text
            continue

        # Detect answer line (e.g. "Answer: B")
        ans_match = re.match(r"(?i)^Answer[:\s]*([a-eA-E])", line)
        if ans_match and current_q:
            current_q["correct"] = ans_match.group(1).upper()
            continue

    # Add the last question if any
    if current_q:
        questions.append(current_q)

    return questions


# -------------------------------
# MAIN PROCESS
# -------------------------------
def process_docx_file(file_path):
    doc = Document(file_path)
    questions = extract_questions(doc)
    total_q = len(questions)
    per_question_score = round(TOTAL_SCORE / total_q, 2) if total_q > 0 else 0

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"

    headers = [
        "S/N", "QUESTION", "OPTION 1", "OPTION 2", "OPTION 3", "OPTION 4", "OPTION 5",
        "CORRECT OPTION", "CORRECT OPTION 2", "SCORE", "DIAGRAM NAME (only jpg)"
    ]
    ws.append(headers)

    for q in questions:
        row = [
            q["sn"],
            q["question"],
            q["options"].get("A", ""),
            q["options"].get("B", ""),
            q["options"].get("C", ""),
            q["options"].get("D", ""),
            q["options"].get("E", ""),
            OPTION_MAP.get(q.get("correct"), ""),
            "",
            per_question_score,
            ""
        ]
        ws.append(row)

    # Save Excel file
    excel_filename = os.path.splitext(os.path.basename(file_path))[0] + ".xlsx"
    excel_path = os.path.join(DONE_FOLDER, excel_filename)
    wb.save(excel_path)

    # Copy processed docx (do not move)
    copy(file_path, os.path.join(PROCESSED_FOLDER, os.path.basename(file_path)))

    return 1


# -------------------------------
# EXECUTION
# -------------------------------
if not os.path.exists(PROCESSED_FOLDER):
    os.makedirs(PROCESSED_FOLDER)
if not os.path.exists(DONE_FOLDER):
    os.makedirs(DONE_FOLDER)

docx_files = [f for f in os.listdir(SOURCE_FOLDER) if f.lower().endswith(".docx")]
total_processed = 0

for docx_file in docx_files:
    file_path = os.path.join(SOURCE_FOLDER, docx_file)
    try:
        total_processed += process_docx_file(file_path)
        print(f"Processed: {docx_file}")
    except Exception as e:
        print(f"Error processing {docx_file}: {e}")

print(f"\nAll done! {total_processed} file(s) processed successfully.")
